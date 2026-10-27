"""服务端 XLSX 导出：openpyxl 生成带表头样式的工作簿。

导出列与导入别名保持一致（日期 / 类别 / 金额 / 备注 / 创建时间），
保证"导出 → 修改 → 再导入"的往返闭环。
"""
import io

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = ("日期", "类别", "金额", "备注", "创建时间")
COLUMN_WIDTHS = (14, 12, 12, 32, 20)
AMOUNT_FORMAT = "0.00"


def build_workbook(expenses) -> bytes:
    """把开销记录写成 xlsx 字节流。"""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "开销记录"
    sheet.append(list(HEADERS))

    header_fill = PatternFill("solid", fgColor="E8EEF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for expense in expenses:
        sheet.append(
            [
                expense.date.strftime("%Y-%m-%d"),
                expense.category,
                float(expense.amount),
                expense.note or "",
                expense.created_at.strftime("%Y-%m-%d %H:%M:%S") if expense.created_at else "",
            ]
        )

    for index, width in enumerate(COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    for row in sheet.iter_rows(min_row=2, min_col=3, max_col=3):
        for cell in row:
            cell.number_format = AMOUNT_FORMAT

    sheet.freeze_panes = "A2"
    if sheet.max_row > 1:
        sheet.auto_filter.ref = sheet.dimensions

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
