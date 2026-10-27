"""开销接口测试：JWT 保护、用户隔离、服务端导入校验。"""
import io
import json
from datetime import date

import openpyxl
from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from rest_framework import status
from rest_framework.test import APITestCase

from .models import Expense

User = get_user_model()


class ExpenseTestBase(APITestCase):
    def setUp(self):
        cache.clear()
        self.user = User.objects.create_user(username="alice", password="secret123")
        self.other = User.objects.create_user(username="bob", password="secret123")

        self.list_url = reverse("expense-list")
        self.import_url = reverse("expense-import")
        self.clear_url = reverse("expense-clear")
        self.token = self._login("alice")
        self.other_token = self._login("bob")

    def _login(self, username):
        response = self.client.post(
            reverse("auth-login"),
            {"username": username, "password": "secret123"},
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        return response.data["access"]

    def auth(self, token=None):
        return {"HTTP_AUTHORIZATION": f"Bearer {token or self.token}"}

    def upload(self, name, content, token=None):
        file = SimpleUploadedFile(name, content)
        return self.client.post(self.import_url, {"file": file}, format="multipart", **self.auth(token))


class ExpenseCrudTests(ExpenseTestBase):
    def test_requires_jwt(self):
        self.assertEqual(self.client.get(self.list_url).status_code, status.HTTP_401_UNAUTHORIZED)

    def test_create_and_list_only_own_records(self):
        response = self.client.post(
            self.list_url,
            {"amount": 35.5, "category": "餐饮", "date": "2026-10-27", "note": "外卖"},
            format="json",
            **self.auth(),
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        expense_id = response.data["id"]

        Expense.objects.create(user=self.other, amount=99, category="购物", date="2026-10-01")

        listing = self.client.get(self.list_url, **self.auth())
        self.assertEqual(listing.status_code, status.HTTP_200_OK)
        self.assertEqual(len(listing.data), 1)
        self.assertEqual(listing.data[0]["id"], expense_id)

    def test_cannot_delete_other_users_record(self):
        other_expense = Expense.objects.create(
            user=self.other, amount=10, category="交通", date="2026-10-01"
        )
        response = self.client.delete(
            reverse("expense-detail", args=[other_expense.id]), **self.auth()
        )
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)
        self.assertTrue(Expense.objects.filter(id=other_expense.id).exists())

    def test_clear_only_affects_current_user(self):
        Expense.objects.create(user=self.user, amount=10, category="餐饮", date="2026-10-01")
        Expense.objects.create(user=self.other, amount=20, category="购物", date="2026-10-02")

        response = self.client.delete(self.clear_url, **self.auth())
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertFalse(Expense.objects.filter(user=self.user).exists())
        self.assertTrue(Expense.objects.filter(user=self.other).exists())


class ImportTests(ExpenseTestBase):
    def _payload(self):
        return {
            "schema_version": 1,
            "expenses": [
                {"amount": 35, "category": "餐饮", "date": "2026-10-27", "note": "外卖"},
                {"amount": 120.5, "category": "购物", "date": "2026-10-26", "note": "超市"},
            ],
        }

    def test_import_json_replaces_existing_records(self):
        Expense.objects.create(user=self.user, amount=1, category="其他", date="2026-01-01")

        response = self.upload("backup.json", json.dumps(self._payload()).encode("utf-8"))
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["imported"], 2)
        self.assertEqual(Expense.objects.filter(user=self.user).count(), 2)
        self.assertFalse(Expense.objects.filter(user=self.user, amount=1).exists())

    def test_import_csv_with_bom_and_quoted_comma(self):
        csv_text = (
            "\ufeffid,date,category,amount,note,created_at\r\n"
            "1,2026-10-27,餐饮,35,\"外卖,加班\",2026-10-27T10:00:00\r\n"
        )
        response = self.upload("backup.csv", csv_text.encode("utf-8"))
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["imported"], 1)
        expense = Expense.objects.get(user=self.user)
        self.assertEqual(expense.note, "外卖,加班")

    def test_import_xlsx_with_chinese_headers(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["日期", "类别", "金额", "备注"])
        sheet.append(["2026-10-27", "餐饮", 42.5, "午餐"])
        sheet.append(["2026-10-26", "交通", 8, "地铁"])
        buffer = io.BytesIO()
        workbook.save(buffer)

        response = self.upload("backup.xlsx", buffer.getvalue())
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["imported"], 2)
        self.assertTrue(Expense.objects.filter(user=self.user, category="交通", amount=8).exists())

    def test_invalid_row_rejects_entire_file_and_keeps_data(self):
        Expense.objects.create(user=self.user, amount=5, category="其他", date="2026-01-01")

        payload = self._payload()
        payload["expenses"][1]["amount"] = -3
        response = self.upload("bad.json", json.dumps(payload).encode("utf-8"))

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("金额", response.data["detail"])
        self.assertEqual(Expense.objects.filter(user=self.user).count(), 1)

    def test_import_requires_jwt(self):
        response = self.client.post(
            self.import_url,
            {"file": SimpleUploadedFile("backup.json", b"[]")},
            format="multipart",
        )
        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_import_rejects_unknown_extension(self):
        response = self.upload("backup.docx", b"not-a-real-doc")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_import_isolated_between_users(self):
        Expense.objects.create(user=self.other, amount=66, category="购物", date="2026-10-01")
        response = self.upload("backup.json", json.dumps(self._payload()).encode("utf-8"))
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        # bob 的数据不应被 alice 的导入影响
        self.assertEqual(Expense.objects.filter(user=self.other).count(), 1)
        self.assertEqual(
            Expense.objects.get(user=self.other).date, date(2026, 10, 1)
        )

    # ---- "伪 Excel"（HTML 表格改名 .xls/.xlsx，账单导出常见） ----

    HTML_TABLE = (
        "<html><head><meta charset='utf-8'></head><body>"
        "<table><tr><th>日期</th><th>类别</th><th>金额</th><th>备注</th></tr>"
        "<tr><td>2026-10-27</td><td>餐饮</td><td>35.5</td><td>外卖,加班</td></tr>"
        "<tr><td>2026-10-26</td><td>交通</td><td>8</td><td></td></tr>"
        "</table></body></html>"
    )

    def test_import_html_table_named_xlsx(self):
        response = self.upload("bill.xlsx", self.HTML_TABLE.encode("utf-8"))
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["imported"], 2)
        self.assertTrue(
            Expense.objects.filter(user=self.user, category="餐饮", amount=35.5).exists()
        )

    def test_import_html_table_named_xls(self):
        response = self.upload("bill.xls", self.HTML_TABLE.encode("utf-8"))
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["imported"], 2)

    def test_import_html_table_gb18030(self):
        response = self.upload("bill.xls", self.HTML_TABLE.encode("gb18030"))
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["imported"], 2)

    def test_import_html_without_table_rejected(self):
        html = "<html><body><p>这是一个网页，没有表格</p></body></html>"
        response = self.upload("bill.xlsx", html.encode("utf-8"))
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("另存为 .xlsx", response.data["detail"])

    def test_import_binary_xls_rejected_with_hint(self):
        ole2 = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64
        response = self.upload("bill.xls", ole2)
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn(".xlsx", response.data["detail"])

    def test_import_html_keeps_data_on_failure(self):
        Expense.objects.create(user=self.user, amount=5, category="其他", date="2026-01-01")
        html = "<html><body><p>no table here</p></body></html>"
        response = self.upload("bill.xls", html.encode("utf-8"))
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(Expense.objects.filter(user=self.user).count(), 1)


class ExportTests(ExpenseTestBase):
    XLSX_CONTENT_TYPE = (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    def setUp(self):
        super().setUp()
        self.export_url = reverse("expense-export")

    def test_requires_jwt(self):
        self.assertEqual(
            self.client.get(self.export_url).status_code, status.HTTP_401_UNAUTHORIZED
        )

    def test_export_returns_xlsx_with_only_own_rows(self):
        Expense.objects.create(
            user=self.user, amount=35.5, category="餐饮", date="2026-10-27", note="外卖"
        )
        Expense.objects.create(
            user=self.other, amount=99, category="购物", date="2026-10-20"
        )

        response = self.client.get(self.export_url, **self.auth())
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response["Content-Type"], self.XLSX_CONTENT_TYPE)
        self.assertIn(".xlsx", response["Content-Disposition"])
        self.assertTrue(response.content.startswith(b"PK"))

        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        rows = list(workbook.active.iter_rows(values_only=True))
        self.assertEqual(list(rows[0]), ["日期", "类别", "金额", "备注", "创建时间"])
        self.assertEqual(len(rows), 2)  # 表头 + 本人 1 条（他人数据不导出）
        self.assertEqual(rows[1][1], "餐饮")
        self.assertAlmostEqual(rows[1][2], 35.5)
        self.assertEqual(rows[1][3], "外卖")

    def test_export_empty_has_header_only(self):
        response = self.client.get(self.export_url, **self.auth())
        rows = list(
            openpyxl.load_workbook(io.BytesIO(response.content)).active.iter_rows(
                values_only=True
            )
        )
        self.assertEqual(len(rows), 1)

    def test_export_then_reimport_round_trip(self):
        Expense.objects.create(
            user=self.user, amount=35, category="餐饮", date="2026-10-27", note="外卖"
        )
        Expense.objects.create(
            user=self.user, amount=120.5, category="购物", date="2026-10-26", note="超市"
        )

        exported = self.client.get(self.export_url, **self.auth())
        upload = SimpleUploadedFile("roundtrip.xlsx", exported.content)
        response = self.client.post(
            self.import_url, {"file": upload}, format="multipart", **self.auth()
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.json()["imported"], 2)
        self.assertEqual(Expense.objects.filter(user=self.user).count(), 2)
        self.assertTrue(
            Expense.objects.filter(user=self.user, category="购物", amount=120.5).exists()
        )
